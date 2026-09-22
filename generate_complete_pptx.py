# -*- coding: utf-8 -*-
"""
Full Python script to generate NirmalNode_Capstone_Defense_NEW.pptx
"""
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

output_pptx = r"c:\Users\ASUS\Downloads\capstone 3.2\NirmalNode_Capstone_Defense_NEW.pptx"
base_dir = r"c:\Users\ASUS\Downloads\capstone 3.2"

def path(p):
    return os.path.join(base_dir, p)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Theme Colors
NAVY = RGBColor(30, 58, 138)        # #1E3A8A
NAVY_LIGHT = RGBColor(37, 99, 235)  # #2563EB
TEAL = RGBColor(15, 118, 110)       # #0F766E
EMERALD = RGBColor(21, 128, 61)     # #15803D
AMBER = RGBColor(180, 83, 9)        # #B45309
ROSE = RGBColor(190, 18, 60)        # #BE123C
DARK_TEXT = RGBColor(15, 23, 42)    # #0F172A
MUTED_TEXT = RGBColor(100, 116, 139)# #64748B
SLATE_BG = RGBColor(248, 250, 252)  # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)
BORDER_COL = RGBColor(226, 232, 240)# #E2E8F0
HIGHLIGHT_ROW = RGBColor(254, 249, 195)
CHAMPION_ROW = RGBColor(224, 242, 254)

def add_header(slide, tag_text, title_text, subtitle_text, category_badge):
    # Top progress bar accent
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = NAVY
    top_bar.line.fill.background()

    # Header text container
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.22), Inches(10.0), Inches(1.15))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = tag_text.upper()
    p0.font.name = 'Arial'
    p0.font.size = Pt(9.5)
    p0.font.bold = True
    p0.font.color.rgb = TEAL

    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.name = 'Arial'
    p1.font.size = Pt(20.5)
    p1.font.bold = True
    p1.font.color.rgb = NAVY

    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = 'Arial'
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED_TEXT

    # Category Badge on Right
    badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.6), Inches(0.32), Inches(2.1), Inches(0.42))
    badge_box.fill.solid()
    badge_box.fill.fore_color.rgb = SLATE_BG
    badge_box.line.color.rgb = BORDER_COL
    badge_box.line.width = Pt(1)
    btf = badge_box.text_frame
    btf.word_wrap = False
    bp = btf.paragraphs[0]
    bp.alignment = PP_ALIGN.CENTER
    bp.text = category_badge
    bp.font.name = 'Arial'
    bp.font.size = Pt(9.5)
    bp.font.bold = True
    bp.font.color.rgb = NAVY

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.4), Inches(12.133), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER_COL
    line.line.fill.background()

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COL, border_width=1):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

print("Helper definitions complete.")
