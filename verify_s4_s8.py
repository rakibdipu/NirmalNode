import pptx

prs = pptx.Presentation(r"c:\Users\ASUS\Downloads\capstone 3.2\NirmalNode_Capstone_Defense_NEW.pptx")

# Slide 4: Check 2026 papers in table
s4 = prs.slides[3]
s4_text = []
for shape in s4.shapes:
    if shape.has_table:
        for r in shape.table.rows:
            s4_text.append(" | ".join(c.text.strip() for c in r.cells))

print("--- Slide 4 Table Rows ---")
papers_2026 = []
for row in s4_text:
    safe_row = row.encode('ascii', 'replace').decode('ascii')
    print(safe_row)
    if "2026" in row:
        papers_2026.append(safe_row)

print(f"\nTotal 2026 rows found: {len(papers_2026)}")
for p in papers_2026:
    print("  *", p)
